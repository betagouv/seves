import io

from django.urls import reverse
import docx
from playwright.sync_api import expect

from core.tests.generic_tests.actions import (
    generic_test_can_cloturer_evenement,
)

from ..models import EvenementAnimal
from .factories import AnalyseFactory, EspeceConcerneeFactory, EvenementAnimalFactory, VeterinaireFactory
from .pages import EvenementAnimalDetailsPage


def test_can_delete_evenement_animal(live_server, page):
    evenement = EvenementAnimalFactory()
    assert EvenementAnimal.objects.count() == 1

    details_page = EvenementAnimalDetailsPage(page, live_server.url)
    details_page.navigate(evenement)
    details_page.delete()
    expect(page.get_by_text(f"L’événement {evenement.numero} a bien été supprimé.")).to_be_visible()

    assert EvenementAnimal.objects.count() == 0
    assert EvenementAnimal._base_manager.get().pk == evenement.pk


def test_can_cloturer_evenement(live_server, page, mocked_authentification_user, mailoutbox):
    evenement = EvenementAnimalFactory(etat=EvenementAnimal.Etat.EN_COURS)
    generic_test_can_cloturer_evenement(
        live_server, page, evenement, mocked_authentification_user, mailoutbox, nav_name="Clôturer l'événement"
    )


def test_can_cloturer_investigation_if_last_remaining_structure(
    live_server, page, mocked_authentification_user, mus_contact
):
    ac_structure = mus_contact.structure
    evenement = EvenementAnimalFactory(etat=EvenementAnimal.Etat.EN_COURS, createur=ac_structure)
    mocked_authentification_user.agent.structure = ac_structure
    evenement.contacts.add(mus_contact)

    details_page = EvenementAnimalDetailsPage(page, live_server.url)
    details_page.navigate(evenement)
    details_page.cloturer(wording="Clôturer l'événement")

    evenement.refresh_from_db()
    assert evenement.etat == EvenementAnimal.Etat.CLOTURE
    assert page.get_by_text("Fin de suivi").count() == 2
    expect(page.get_by_text(f"L'événement n°{evenement.numero} a bien été clôturé.")).to_be_visible()


def test_can_download_document_evenement_animal(live_server, page):
    evenement = EvenementAnimalFactory(date_publication=None)

    details_page = EvenementAnimalDetailsPage(page, live_server.url)
    details_page.navigate(evenement)
    download = details_page.download().value
    assert download.suggested_filename == f"evenement_animal_{evenement.numero}.docx"


def test_document_evenement_animal_contains_sub_objects(client):
    evenement = EvenementAnimalFactory()
    analyse = AnalyseFactory(evenement=evenement)
    veterinaire = VeterinaireFactory(evenement=evenement)
    espece_concernee = EspeceConcerneeFactory(evenement=evenement, identifiant="ID-ESPECE-123")

    response = client.post(reverse("sa:evenement-animal-export-document", kwargs={"pk": evenement.pk}))

    assert response.status_code == 200
    document = docx.Document(io.BytesIO(response.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    tables_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert evenement.numero in text
    assert evenement.maladie.name in text
    assert analyse.laboratoire.name in text
    assert veterinaire.nom_structure in text
    assert espece_concernee.identifiant in tables_text
    assert "{{" not in text + tables_text
    assert "{%" not in text + tables_text
